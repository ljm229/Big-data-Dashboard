"""Read four source1 workbooks into the risk module's snapshot; never modify Excel.

Requires Python 3 + openpyxl. No collection, network, database or scheduled job.
Conflicting daily keys / missing schemas fail before replacing the previous snapshot.
"""
from datetime import date, datetime, timezone
from pathlib import Path
import hashlib
import json
import math
import re
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "数据源1"
OUTPUT = ROOT / "web/src/data/riskData.json"


def norm(value):
    return re.sub(r"\s+", "", str(value or "")).replace("（", "(").replace("）", ")")


def iso(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    if re.fullmatch(r"\d{8}", text):
        text = f"{text[:4]}-{text[4:6]}-{text[6:]}"
    if re.match(r"^\d{4}-\d{2}-\d{2}(?:$|[ T])", text):
        try:
            return date.fromisoformat(text[:10]).isoformat()
        except ValueError:
            pass
    return ""


def number(value):
    if value is None or isinstance(value, bool):
        return None
    text = str(value).strip().replace(",", "")
    try:
        result = float(text[:-1]) / 100 if text.endswith("%") else float(text)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def load(folder, pattern, required):
    paths = [p for p in (SOURCE / folder).glob(pattern) if not p.name.startswith("~$")]
    if not paths:
        raise ValueError(f"缺少源文件: {folder}/{pattern}")
    path = max(paths, key=lambda p: (p.stat().st_mtime_ns, p.name))
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["data"] if "data" in wb.sheetnames else wb.worksheets[0]
        ws.reset_dimensions()  # Some platform exports incorrectly declare A1:A1.
        iterator = ws.iter_rows(values_only=True)
        headers = [str(v or "").strip() for v in next(iterator)]
        if set(required) - set(headers):
            raise ValueError(f"{path.name} 缺少列: {sorted(set(required) - set(headers))}")
        rows = [(i, dict(zip(headers, values))) for i, values in enumerate(iterator, 2)]
        metadata = {"path": path.relative_to(SOURCE).as_posix(), "sheet": ws.title,
                    "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        return metadata, rows
    finally:
        wb.close()


def unique(rows, keys):
    result = {}
    duplicates = 0
    for row in rows:
        key = tuple(row[k] for k in keys)
        if key in result:
            old = {k: v for k, v in result[key].items() if k != "row"}
            new = {k: v for k, v in row.items() if k != "row"}
            if old != new:
                raise ValueError(f"源数据主键冲突，未覆盖输出: {key}")
            duplicates += 1
        else:
            result[key] = row
    return sorted(result.values(), key=lambda r: tuple(str(r[k]) for k in keys)), duplicates


def main():
    files, stats = {}, {}
    files["stores"], rows = load("翱象", "*城市门店及上线进度*.xlsx", ["门店", "城市", "是否上线"])
    stores = [{"name": norm(r["门店"]), "city": str(r["城市"] or "").strip(),
               "status": str(r["是否上线"] or "").strip(), "row": i}
              for i, r in rows if norm(r["门店"]) and str(r["城市"] or "").strip()]
    stores, stats["storeDuplicates"] = unique(stores, ["name"])

    fact_fields = {"profit": "预计毛利(含平台后返)", "orders": "有效订单量", "paid": "有效订单金额（实付）",
                   "refundRate": "退款率", "refundOrders": "退款订单量", "refundAmount": "退款金额",
                   "negativeOrders": "负毛利订单量", "negativeOrderRate": "负毛利订单占比"}
    files["facts"], rows = load("翱象", "*渠道门店周期数据*.xlsx", ["日期", "渠道", "门店", *fact_fields.values()])
    facts = []
    for i, r in rows:
        day, store, channel = iso(r["日期"]), norm(r["门店"]), str(r["渠道"] or "").strip()
        metrics = {k: number(r.get(v)) for k, v in fact_fields.items()}
        if day and store and channel and any(v is not None for v in metrics.values()):
            facts.append({"date": day, "store": store, "channel": channel, **metrics, "row": i})
    facts, stats["factDuplicates"] = unique(facts, ["store", "channel", "date"])

    supply_fields = {"onShelf": "在架商品数", "stockout": "缺货商品数", "attendance": "商品出勤率",
                     "absent": "缺勤商品数", "absentLoss": "缺勤商品损失金额"}
    files["supply"], rows = load("淘宝闪购商家", "*按店铺汇总*.xlsx", ["日期", "门店名称", *supply_fields.values()])
    supply = []
    for i, r in rows:
        day, store = iso(r["日期"]), norm(r["门店名称"])
        metrics = {k: number(r.get(v)) for k, v in supply_fields.items()}
        if day and store and any(v is not None for v in metrics.values()):
            supply.append({"date": day, "store": store, **metrics, "row": i})
    supply, stats["supplyDuplicates"] = unique(supply, ["store", "date"])

    quality_fields = {"sellout": "动销商品售罄率", "pickingError": "错漏拣率", "warehouseT": "仓T",
                      "imReply": "IM 3分钟回复率", "merchantIssue": "商责问题订单率", "rating": "店铺分"}
    files["quality"], rows = load("翱象", "*门店营运考核指标*.xlsx", ["日期", "门店名称", *quality_fields.values()])
    quality = []
    for i, r in rows:
        day, store = iso(r["日期"]), norm(r["门店名称"])
        metrics = {k: number(r.get(v)) for k, v in quality_fields.items()}
        if day and store and store != "门店名称" and any(v is not None for v in metrics.values()):
            quality.append({"date": day, "store": store, **metrics, "row": i})
    stats["qualitySkippedRows"] = len(rows) - len(quality)
    quality, stats["qualityDuplicates"] = unique(quality, ["store", "date"])
    if not all([stores, facts, supply, quality]):
        raise ValueError("存在空数据源，保留原输出，请核对源文件")
    payload = {"generatedAt": datetime.now(timezone.utc).isoformat(), "files": files, "stats": stats,
               "stores": stores, "facts": facts, "supply": supply, "quality": quality}
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT.with_suffix(".json.tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":"), allow_nan=False), encoding="utf-8")
    temporary.replace(OUTPUT)
    print(json.dumps({"output": str(OUTPUT), "counts": {k: len(payload[k]) for k in ["stores", "facts", "supply", "quality"]}, "stats": stats}, ensure_ascii=False))


if __name__ == "__main__":
    main()
