"""Fast streaming import for 毛利订单 (55MB+) using openpyxl read-only mode.
Same aggregation口径 as scripts/sync-order-margin.mjs. Read-only, never writes back to Excel."""
import hashlib, json, sys
from collections import defaultdict
from datetime import datetime, date
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = ROOT / "数据源" / "翱象"
OUT = ROOT / "web" / "web" / "src" / "data" / "orderMarginData.json"
REASON_KEYS = ['商品毛利为负', '营销折扣过高', '配送成本过高', '平台费用占比高', '其他']

def norm(s):
    import re
    s = str(s or '')
    s = re.sub(r'\s+', '', s).replace('（', '(').replace('）', ')')
    return s

def channel_of(s):
    t = str(s or '').strip()
    return 'POS渠道' if t == 'POS' else t

def iso(v):
    if isinstance(v, datetime): return v.date().isoformat()
    if isinstance(v, date): return v.isoformat()
    t = str(v or '').strip()[:10]
    if len(t) == 10 and t[4] == '-' and t[7] == '-': return t
    t8 = str(v or '').strip()
    if len(t8) == 8 and t8.isdigit(): return f'{t8[:4]}-{t8[4:6]}-{t8[6:8]}'
    # datetime string like 2026-09-19 12:00:00
    try:
        return datetime.fromisoformat(str(v).strip().replace('/', '-')).date().isoformat()
    except Exception:
        return ''

def num(v):
    if v is None or v == '': return None
    try:
        n = float(str(v).replace(',', '').strip())
        return n
    except Exception:
        return None

def n0(v):
    n = num(v)
    return n if n is not None else 0.0

def money(v): return round(v * 100) / 100

def classify(row):
    product = n0(row.get('应收商品总额')) + n0(row.get('商家商品优惠')) + n0(row.get('商品采购成本'))
    marketing = n0(row.get('商家商品优惠')) + n0(row.get('商家整单优惠'))
    delivery = n0(row.get('应收配送费')) + n0(row.get('配送费优惠')) + n0(row.get('商家自配送成本')) + n0(row.get('平台配送服务费'))
    platform = n0(row.get('佣金')) + n0(row.get('其他平台费')) + n0(row.get('平台配送服务费'))
    parts = sorted([('商品毛利为负', product), ('营销折扣过高', marketing), ('配送成本过高', delivery), ('平台费用占比高', platform)], key=lambda x: x[1])
    return parts[0][0] if parts[0][1] < 0 else '其他'

def main():
    files = [p for p in SRC_DIR.iterdir() if p.suffix == '.xlsx' and ('订单毛利' in p.name or '毛利订单' in p.name) and not p.name.startswith('~$')]
    if not files: sys.exit('缺少源文件: 翱象/*订单毛利*.xlsx')
    src = max(files, key=lambda p: (p.stat().st_mtime, p.name))
    sha = hashlib.sha256(src.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = next((wb[s] for s in wb.sheetnames if '订单毛利' in s), wb[wb.sheetnames[0]])
    rows = ws.iter_rows(values_only=True)
    header = [str(c or '').strip() for c in next(rows)]
    idx = {h: i for i, h in enumerate(header)}
    for k in ['门店名称', '渠道名称', '创建时间', '预计毛利']:
        if k not in idx: sys.exit(f'{src.name} 缺少列: {k}')
    def g(r, k): return r[idx[k]] if k in idx else None
    facts = {}
    raw = skipped = pos = neg = refund = 0
    ch_tot = defaultdict(int); ch_neg = defaultdict(int)
    for r in rows:
        raw += 1
        dt = g(r, '创建时间') or g(r, '订单完成时间')
        d = iso(dt); store = norm(g(r, '门店名称')); ch = channel_of(g(r, '渠道名称')); profit = num(g(r, '预计毛利'))
        if not d or not store or not ch or profit is None:
            skipped += 1; continue
        key = (d, store, ch)
        f = facts.get(key)
        if f is None:
            f = dict(date=d, store=store, storeCode=str(g(r, '门店编码') or '').strip(), channel=ch,
                     orders=0, negOrders=0, negGt3=0, loss=0.0, revenue=0.0, profitSum=0.0,
                     marketing=0.0, delivery=0.0, platform=0.0, refundOrders=0, refundProfit=0.0,
                     reasons={k: {'count': 0, 'amount': 0.0} for k in REASON_KEYS})
            facts[key] = f
        rd = {h: g(r, h) for h in header}
        f['orders'] += 1; f['revenue'] += n0(g(r, '应收')); f['profitSum'] += profit
        f['marketing'] += n0(g(r, '商家商品优惠')) + n0(g(r, '商家整单优惠')) + n0(g(r, '配送费优惠'))
        f['delivery'] += n0(g(r, '应收配送费')) + n0(g(r, '商家自配送成本')) + n0(g(r, '平台配送服务费'))
        f['platform'] += n0(g(r, '佣金')) + n0(g(r, '其他平台费'))
        ch_tot[ch] += 1
        rv = g(r, '退款单号')
        if rv not in (None, '') and str(rv).strip() != '':
            refund += 1; f['refundOrders'] += 1; f['refundProfit'] += profit
        if profit < 0:
            neg += 1; f['negOrders'] += 1; f['loss'] += profit
            if profit <= -3: f['negGt3'] += 1
            rs = classify(rd); f['reasons'][rs]['count'] += 1; f['reasons'][rs]['amount'] += profit
            ch_neg[ch] += 1
        else: pos += 1
        if raw % 100000 == 0: print(f'...{raw}', flush=True)
    fl = sorted(facts.values(), key=lambda f: (f['date'], f['store'], f['channel']))
    for f in fl:
        for k in ['loss', 'revenue', 'profitSum', 'marketing', 'delivery', 'platform', 'refundProfit']:
            f[k] = money(f[k])
        for k in REASON_KEYS: f['reasons'][k]['amount'] = money(f['reasons'][k]['amount'])
    dates = sorted({f['date'] for f in fl})
    payload = dict(generatedAt=datetime.now().astimezone().isoformat(), source=dict(
        path=str(src.relative_to(ROOT)).replace('\\', '/'), sheet=ws.title, sha256=sha,
        note='全量订单；负毛利结构按预计毛利<0；＞3元=预计毛利≤-3；变化桥字段=应收/预计毛利/营销/配送/平台费/退款单毛利。'),
        reasonKeys=REASON_KEYS, stats=dict(rawRows=raw, skipped=skipped, positiveOrders=pos, negativeOrders=neg,
        refundOrders=refund, factRows=len(fl), orderRows=sum(f['orders'] for f in fl),
        negOrderRows=sum(f['negOrders'] for f in fl), negGt3=sum(f['negGt3'] for f in fl),
        loss=money(sum(f['loss'] for f in fl)), revenue=money(sum(f['revenue'] for f in fl)),
        profitSum=money(sum(f['profitSum'] for f in fl)), dateFrom=dates[0] if dates else None,
        dateTo=dates[-1] if dates else None, channels=dict(ch_tot), channelNeg=dict(ch_neg)), facts=fl)
    tmp = OUT.with_suffix('.json.tmp')
    tmp.write_text(json.dumps(payload), encoding='utf-8')
    tmp.replace(OUT)
    print(json.dumps(dict(ok=True, file=src.name, rawRows=raw, factRows=len(fl), dateFrom=payload['stats']['dateFrom'], dateTo=payload['stats']['dateTo']), ensure_ascii=False))

if __name__ == '__main__':
    main()
